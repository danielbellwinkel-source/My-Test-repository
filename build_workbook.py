# -*- coding: utf-8 -*-
"""Baut Training_2.xlsx mit dem B6-Trainingsblock (Hand Balancing / Calisthenics).
Standalone-Variante: keine Training_1.xlsx vorhanden -> Datei wird neu erstellt.
Alle Auswertungen per Formel, Excel-Tabellen (ListObjects), Dropdowns, bedingte Formatierung.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.worksheet.table import Table, TableStyleInfo, TableColumn, TableFormula
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
import datetime

# ----------------------------------------------------------------------------
# Stil-Helfer (durchgaengig Arial)
# ----------------------------------------------------------------------------
FONT = "Arial"
C_DARK   = "1F3864"   # Titel-/Kopf-Hintergrund (dunkelblau)
C_MID    = "2E5496"
C_HEAD   = "305496"   # Tabellenkopf
C_SUB    = "D9E1F2"   # Sektions-/Subkopf hell
C_BAND   = "F2F6FC"   # Zebra hell
C_BLOCK  = "FFE699"   # Block-Variable
C_GREEN  = "C6EFCE"
C_GREENT = "006100"
C_YELLOW = "FFEB9C"
C_ORANGE = "FFD580"
C_RED    = "FFC7CE"
C_REDT   = "9C0006"

def F(size=10, bold=False, color="000000", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)

def fill(hex_):
    return PatternFill("solid", fgColor=hex_)

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="top", wrap_text=True)

def title(ws, cell, text, span=None, size=14):
    c = ws[cell]
    c.value = text
    c.font = F(size, bold=True, color="FFFFFF")
    c.fill = fill(C_DARK)
    c.alignment = Alignment(horizontal="left", vertical="center")
    if span:
        ws.merge_cells(span)
        for row in ws[span]:
            for cc in row:
                cc.fill = fill(C_DARK)

def section(ws, cell, text, span=None):
    c = ws[cell]
    c.value = text
    c.font = F(11, bold=True, color="1F3864")
    c.fill = fill(C_SUB)
    c.alignment = Alignment(horizontal="left", vertical="center")
    if span:
        ws.merge_cells(span)
        for row in ws[span]:
            for cc in row:
                cc.fill = fill(C_SUB)

def header_row(ws, row, headers, start_col=1):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + i, value=h)
        c.font = F(10, bold=True, color="FFFFFF")
        c.fill = fill(C_HEAD)
        c.alignment = CENTER
        c.border = BORDER

def put(ws, row, col, value, bold=False, align=LEFT, border=True, size=10, color="000000", numfmt=None, fillhex=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = F(size, bold=bold, color=color)
    c.alignment = align
    if border:
        c.border = BORDER
    if numfmt:
        c.number_format = numfmt
    if fillhex:
        c.fill = fill(fillhex)
    return c

wb = Workbook()
wb.remove(wb.active)

# Datum (heute) fuer Demo-Zeile / Defaults
TODAY = datetime.date(2026, 6, 8)

# ============================================================================
# 1) B6_Exercise_Library  (zuerst, weil andere Sheets darauf referenzieren)
# ============================================================================
lib = wb.create_sheet("B6_Exercise_Library")
title(lib, "A1", "B6 — Exercise Library (Referenz fuer Voice-Logging)", "A1:H1")
lib["A2"] = ("Aliase Komma-getrennt, deutsch + englisch. Das Logging-Modell mappt gesprochene Begriffe "
             "ueber diese Tabelle auf die ExerciseID. TracksLoad/Hold/Side = J/N.")
lib["A2"].font = F(9, italic=True); lib.merge_cells("A2:H2"); lib["A2"].alignment = WRAP

LIB_HEADERS = ["ExerciseID","Aliases","DisplayName","Category","MuscleTag","TracksLoad","TracksHold","TracksSide"]
LIB_HEADER_ROW = 3
header_row(lib, LIB_HEADER_ROW, LIB_HEADERS)

# (ID, Aliases, DisplayName, Category, MuscleTag, Load, Hold, Side)
LIBRARY = [
 ("90","neunzig grad,90,90 grad,90 grad liegestuetz,ninety,90 degree push up,neunzig","90 Degree Push-up","Skill90","Brust/Schulter/Planche","N","N","N"),
 ("OAHS_line","oahs,oahs line,einarmiger handstand,one arm handstand,weight shift,gewichtsverlagerung,oahs holds","OAHS Line / Weight-Shift / Holds","SkillOAHS","Schulter/Core/Balance","N","J","J"),
 ("OAHS_shapes","oahs shapes,straddle,half straddle,diamond,legs together,oahs formen,shape density","OAHS Shape-Density","SkillOAHS","Schulter/Core/Balance","N","J","J"),
 ("HS_line","handstand,handstand line,hs line,modern line,balance,handstand balance,linie","Handstand Line / Balance","SkillOAHS","Schulter/Core/Balance","N","J","N"),
 ("OAHS_flag_supported","supported flag,pre flag,supported pre flag,flag supported,unterstuetzter flag,pre flag lean","Supported Pre-Flag Lean","SkillFlag","Schulter/Side-Body","N","J","J"),
 ("flag_sideline","sideline,sideline hold,mini flag,flag hold,flag,seitlicher flag","Sideline / Mini-Flag Hold","SkillFlag","Schulter/Side-Body","N","J","J"),
 ("pseudo_planche_pushup","pseudo planche,pseudo planche push up,ppp,pseudo planche liegestuetz","Pseudo-Planche Push-up","Planche","Schulter/Brust","N","N","N"),
 ("planche_lean","planche lean,lean,planche neigung,plankenlean","Planche Lean","Planche","Schulter/Core","N","J","N"),
 ("planche_hold","planche hold,tuck planche,advanced tuck,adv tuck,straddle planche,planche halten","Planche Hold (Tuck/Adv/Straddle)","Planche","Schulter/Core","N","J","N"),
 ("WPU","weighted pull up,wpu,gewichtetes klimmziehen,klimmzug mit gewicht,weighted pullup","Weighted Pull-up","Pull","Lat/Biceps","J","N","N"),
 ("pullup","pull up,pullup,klimmzug,klimmziehen","Pull-up","Pull","Lat/Biceps","J","N","N"),
 ("chin_up","chin up,chinup,kinnzug,unterhand klimmzug","Chin-up","Pull","Lat/Biceps","J","N","N"),
 ("lat_pulldown","lat pulldown,pulldown,latzug,latziehen","Lat Pulldown","Pull","Lat","J","N","N"),
 ("chest_row","chest supported row,seal row,brustgestuetztes rudern,seal row,rudern gestuetzt","Chest-Supported / Seal Row","Pull","Oberer Ruecken","J","N","N"),
 ("oa_row","one arm row,einarmiges rudern,oa row,unilateral row","One-Arm Row","Pull","Oberer Ruecken","J","N","J"),
 ("rear_delt_fly","rear delt fly,reverse fly,hintere schulter,rear delt","Rear-Delt Fly","Pull","Hintere Schulter","J","N","N"),
 ("face_pull","face pull,facepull,gesichtszug","Face Pull","Pull","Hintere Schulter/Rotatoren","J","N","N"),
 ("scap_serratus","scap,serratus,scapula,schulterblatt,scap pull,serratus push","Scap / Serratus","Mobility","Schulterblatt","N","N","N"),
 ("ext_rotation","external rotation,aussenrotation,ar,rotatorenmanschette,cuff","External Rotation","Mobility","Rotatorenmanschette","J","N","J"),
 ("wall_hspu","wall hspu,hspu,handstand push up,handstand liegestuetz,wand hspu","Wall HSPU","Push","Schulter/Trizeps","J","N","N"),
 ("pike_press","pike press,pike push up,db press,schulterdruecken,overhead press,ohp,pike","Pike / DB Vertical Press","Push","Schulter","J","N","N"),
 ("ring_dips","ring dips,dips,barren,ringdips","Ring Dips / Dips","Push","Brust/Trizeps","J","N","N"),
 ("deficit_pushup","deficit push up,deficit pushup,liegestuetz deficit,push up,liegestuetz","Deficit Push-up","Push","Brust/Trizeps","N","N","N"),
 ("lateral_raise","lateral raise,seitheben,lat raise,seitliches heben","Lateral Raise","Push","Seitliche Schulter","J","N","N"),
 ("triceps","triceps,trizeps,trizepsdruecken,pushdown","Triceps","Push","Trizeps","J","N","N"),
 ("biceps","biceps,bizeps,curl,bizeps curl","Biceps","Pull","Bizeps","J","N","N"),
 ("bss","bulgarian,bulgarian split squat,bss,split squat,ausfallschritt,ffe split squat,front foot elevated","Bulgarian / FFE Split Squat","Legs","Quad/Glute","J","N","J"),
 ("leg_press","leg press,beinpresse,hack squat,hack,belt squat,guertelkniebeuge","Leg Press / Hack / Belt Squat","Legs","Quad/Glute","J","N","N"),
 ("ham_curl","hamstring curl,beinbeuger,leg curl,beincurl","Hamstring Curl","Legs","Hamstrings","J","N","N"),
 ("hip_thrust","hip thrust,hueftstoss,glute bridge,hueftheben","Hip Thrust","Legs","Glute","J","N","N"),
 ("back_ext_iso","back extension,rueckenstrecker,back ext iso,hyperextension,iso hold ruecken","Back-Extension Iso","Legs","Posteriore Kette/LWS","N","J","N"),
 ("calves","calves,waden,wadenheben,calf raise","Calves","Legs","Waden","J","N","N"),
 ("hinge_pattern","hinge,hinge pattern,dowel hinge,hueftbeuge,rdl pattern,deadlift pattern,scharnier","Hinge Pattern (Dowel/Rebuild)","Legs","Posteriore Kette/LWS","J","N","N"),
 ("pallof","pallof,pallof press,anti rotation","Pallof Press","Core","Anti-Rotation","J","N","J"),
 ("dead_bug","dead bug,deadbug,toter kaefer","Dead Bug","Core","Anti-Extension","N","N","N"),
 ("side_plank","side plank,seitstuetz,seitlicher unterarmstuetz","Side Plank","Core","Lateral/Anti-Lateral","N","J","J"),
 ("hollow","hollow,hollow hold,hollow body,compression,kompression,hohlkreuz hold","Hollow / Compression Core","Core","Anteriore Kette","N","J","N"),
 ("bws_mobility","bws,bws mobility,brustwirbelsaeule,thoracic,t spine,extension rotation,ribs down,atmung,downshift","BWS / Thoracic Mobility","Mobility","BWS/Atmung","N","J","N"),
 ("wrist_prep","wrist prep,handgelenk,handgelenk prep,handgelenke,wrist warm up,wrist","Wrist Prep","Wrist","Handgelenke","N","N","N"),
 ("wrist_strength","wrist strength,handgelenk kraft,wrist curl,handgelenk kraeftigung","Wrist Strengthening","Wrist","Handgelenke","J","N","N"),
]
r = LIB_HEADER_ROW + 1
for row in LIBRARY:
    for i, val in enumerate(row):
        align = LEFT if i in (1,2,3,4) else CENTER
        put(lib, r, 1+i, val, align=align)
    if (r % 2) == 0:
        for i in range(len(row)):
            if not lib.cell(r,1+i).fill or lib.cell(r,1+i).fill.fgColor.rgb in (None,"00000000"):
                lib.cell(r,1+i).fill = fill(C_BAND)
    r += 1
LIB_LAST = r - 1

# Hilfsliste DayTypes (fuer benannten Bereich + Dropdown) in Spalte K
put(lib, LIB_HEADER_ROW, 11, "DayTypes", bold=True, align=CENTER, fillhex=C_HEAD)
lib.cell(LIB_HEADER_ROW,11).font = F(10, bold=True, color="FFFFFF")
DAYTYPES = ["T1","T2","T3","T4","T5","Rest"]
for i, d in enumerate(DAYTYPES):
    put(lib, LIB_HEADER_ROW+1+i, 11, d, align=CENTER)
DT_LAST = LIB_HEADER_ROW+len(DAYTYPES)

# ListObject fuer Library
lib_tab = Table(displayName="tblLib", ref=f"A{LIB_HEADER_ROW}:H{LIB_LAST}")
lib_tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True,
                                        showFirstColumn=False, showLastColumn=False, showColumnStripes=False)
lib.add_table(lib_tab)

widths = {"A":22,"B":46,"C":30,"D":12,"E":24,"F":11,"G":11,"H":11,"I":3,"J":3,"K":12}
for col,w in widths.items():
    lib.column_dimensions[col].width = w
lib.freeze_panes = "A4"

# Benannte Bereiche
wb.defined_names.add(DefinedName("ExerciseIDs", attr_text=f"B6_Exercise_Library!$A${LIB_HEADER_ROW+1}:$A${LIB_LAST}"))
wb.defined_names.add(DefinedName("DayTypes", attr_text=f"B6_Exercise_Library!$K${LIB_HEADER_ROW+1}:$K${DT_LAST}"))

# ============================================================================
# 2) B6_Training_Log  (primaere Eingabeflaeche, Excel-Tabelle, append-only)
# ============================================================================
log = wb.create_sheet("B6_Training_Log")
LOG_HEADERS = ["Date","DayType","Week","ExerciseID","Category","Side","Sets","Reps_or_Scheme",
               "Load_kg","TopRPE","Pain_0_10","Quality_0_5","BestHold_s","CleanReps","Notes","Key"]
HR = 1
header_row(log, HR, LOG_HEADERS)

# Demo-Zeile (Zeile 2)
demo_row = 2
log.cell(demo_row,1, TODAY); log.cell(demo_row,1).number_format = "yyyy-mm-dd"
log.cell(demo_row,2,"T1")
log.cell(demo_row,3,1)
log.cell(demo_row,4,"90")
log.cell(demo_row,5,'=IFERROR(XLOOKUP([@ExerciseID],tblLib[ExerciseID],tblLib[Category]),"")')
log.cell(demo_row,6,"NA")
log.cell(demo_row,7,5)
log.cell(demo_row,8,"1")
log.cell(demo_row,9,0)
log.cell(demo_row,10,8.5)
log.cell(demo_row,11,1)
log.cell(demo_row,12,None)
log.cell(demo_row,13,None)
log.cell(demo_row,14,5)
log.cell(demo_row,15,"Demo-Zeile: 90 Grad schwer, je 1 saubere Rep, Handgelenk gut")
log.cell(demo_row,16,'=[@DayType]&"_W"&[@Week]')
for col in range(1,17):
    cc = log.cell(demo_row, col)
    cc.font = F(10)
    cc.border = BORDER
    cc.alignment = LEFT if col in (8,15) else CENTER

log_tab = Table(displayName="tblLog", ref=f"A{HR}:P{demo_row}")
log_tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
# Explizite Spalten: Category & Key als berechnete Spalten -> fuellen sich bei angehaengten Zeilen
cols = []
for i, h in enumerate(LOG_HEADERS):
    tc = TableColumn(id=i + 1, name=h)
    if h == "Category":
        tc.calculatedColumnFormula = TableFormula(
            attr_text='IFERROR(XLOOKUP([@ExerciseID],tblLib[ExerciseID],tblLib[Category]),"")')
    elif h == "Key":
        tc.calculatedColumnFormula = TableFormula(attr_text='[@DayType]&"_W"&[@Week]')
    cols.append(tc)
log_tab.tableColumns = cols
log_tab.autoFilter = AutoFilter(ref=f"A{HR}:P{demo_row}")  # noetig, da manuelle tableColumns
log.add_table(log_tab)

logw = {"A":12,"B":9,"C":7,"D":20,"E":12,"F":7,"G":7,"H":15,"I":9,"J":8,"K":10,"L":11,"M":11,"N":10,"O":42,"P":10}
for col,w in logw.items():
    log.column_dimensions[col].width = w
log.freeze_panes = "A2"

# Datenvalidierung
dv_day = DataValidation(type="list", formula1="DayTypes", allow_blank=False, showDropDown=False)
dv_ex  = DataValidation(type="list", formula1="ExerciseIDs", allow_blank=False, showDropDown=False)
dv_side= DataValidation(type="list", formula1='"L,R,NA"', allow_blank=True, showDropDown=False)
dv_week= DataValidation(type="whole", operator="between", formula1="1", formula2="4", allow_blank=True)
dv_rpe = DataValidation(type="decimal", operator="between", formula1="0", formula2="10", allow_blank=True)
dv_pain= DataValidation(type="whole", operator="between", formula1="0", formula2="10", allow_blank=True)
dv_qual= DataValidation(type="whole", operator="between", formula1="0", formula2="5", allow_blank=True)
RNG = 1000  # Vorrat fuer angehaengte Zeilen
for dv, col in [(dv_day,"B"),(dv_week,"C"),(dv_ex,"D"),(dv_side,"F"),(dv_rpe,"J"),(dv_pain,"K"),(dv_qual,"L")]:
    dv.add(f"{col}2:{col}{RNG}")
    log.add_data_validation(dv)

# Bedingte Formatierung
pain_range = f"K2:K{RNG}"
log.conditional_formatting.add(pain_range, CellIsRule(operator="greaterThanOrEqual", formula=["4"],
        fill=fill(C_RED), font=F(10,color=C_REDT)))
log.conditional_formatting.add(pain_range, CellIsRule(operator="equal", formula=["3"], fill=fill(C_ORANGE)))
log.conditional_formatting.add(pain_range, CellIsRule(operator="between", formula=["0","2"],
        fill=fill(C_GREEN), font=F(10,color=C_GREENT)))
log.conditional_formatting.add(f"J2:J{RNG}", CellIsRule(operator="greaterThan", formula=["8.5"],
        fill=fill(C_YELLOW)))

# ============================================================================
# 3) B6_Weekly_Summary  (Auto-Rollups per Formel)
# ============================================================================
ws_sum = wb.create_sheet("B6_Weekly_Summary")
title(ws_sum, "A1", "B6 — Weekly Summary (automatisch aus B6_Training_Log)", "A1:F1")
ws_sum["A2"] = "Alle Werte per Formel ueber strukturierte Referenzen. Angehaengte Log-Zeilen fliessen automatisch ein."
ws_sum["A2"].font = F(9, italic=True); ws_sum.merge_cells("A2:F2")

SUM_HR = 3
sum_headers = ["Metrik","W1","W2","W3","W4","Block gesamt"]
header_row(ws_sum, SUM_HR, sum_headers)

# Zeilendefinitionen: (Label, formel-template fuer Woche w, formel fuer gesamt, numfmt)
def pushf(scope):
    # scope: 'tblLog[Week],{w}' oder '' fuer gesamt
    base = "SUMIFS(tblLog[Sets],tblLog[Category],\"{cat}\"{flt})"
    crit = scope
    parts = []
    for cat in ("Skill90","Planche","Push"):
        parts.append(base.format(cat=cat, flt=crit))
    return "=" + "+".join(parts)

metrics = []
def m(label, wtmpl, gtmpl, numfmt="0.0"):
    metrics.append((label, wtmpl, gtmpl, numfmt))

# Push (Skill90+Planche+Push)
m("Push — effektive Saetze (90/Planche/Push)",
  lambda w: pushf(f",tblLog[Week],{w}"),
  pushf(""), "0.0")
m("Pull — effektive Saetze",
  lambda w: f'=SUMIFS(tblLog[Sets],tblLog[Category],"Pull",tblLog[Week],{w})',
  '=SUMIFS(tblLog[Sets],tblLog[Category],"Pull")', "0.0")
m("Legs — effektive Saetze",
  lambda w: f'=SUMIFS(tblLog[Sets],tblLog[Category],"Legs",tblLog[Week],{w})',
  '=SUMIFS(tblLog[Sets],tblLog[Category],"Legs")', "0.0")
m("90 — Touches (Eintraege)",
  lambda w: f'=COUNTIFS(tblLog[ExerciseID],"90",tblLog[Week],{w})',
  '=COUNTIFS(tblLog[ExerciseID],"90")', "0")
m("90 — saubere Reps gesamt",
  lambda w: f'=SUMIFS(tblLog[CleanReps],tblLog[ExerciseID],"90",tblLog[Week],{w})',
  '=SUMIFS(tblLog[CleanReps],tblLog[ExerciseID],"90")', "0")
m("OAHS — Skill-Bloecke (Eintraege)",
  lambda w: f'=COUNTIFS(tblLog[Category],"SkillOAHS",tblLog[Week],{w})',
  '=COUNTIFS(tblLog[Category],"SkillOAHS")', "0")
m("Planche — Dosen (Eintraege)",
  lambda w: f'=COUNTIFS(tblLog[Category],"Planche",tblLog[Week],{w})',
  '=COUNTIFS(tblLog[Category],"Planche")', "0")
m("OAHS Best-Hold links (s)",
  lambda w: f'=MAXIFS(tblLog[BestHold_s],tblLog[Category],"SkillOAHS",tblLog[Side],"L",tblLog[Week],{w})',
  '=MAXIFS(tblLog[BestHold_s],tblLog[Category],"SkillOAHS",tblLog[Side],"L")', "0")
m("OAHS Best-Hold rechts (s)",
  lambda w: f'=MAXIFS(tblLog[BestHold_s],tblLog[Category],"SkillOAHS",tblLog[Side],"R",tblLog[Week],{w})',
  '=MAXIFS(tblLog[BestHold_s],tblLog[Category],"SkillOAHS",tblLog[Side],"R")', "0")
m("Flag Best-Hold links (s)",
  lambda w: f'=MAXIFS(tblLog[BestHold_s],tblLog[Category],"SkillFlag",tblLog[Side],"L",tblLog[Week],{w})',
  '=MAXIFS(tblLog[BestHold_s],tblLog[Category],"SkillFlag",tblLog[Side],"L")', "0")
m("Flag Best-Hold rechts (s)",
  lambda w: f'=MAXIFS(tblLog[BestHold_s],tblLog[Category],"SkillFlag",tblLog[Side],"R",tblLog[Week],{w})',
  '=MAXIFS(tblLog[BestHold_s],tblLog[Category],"SkillFlag",tblLog[Side],"R")', "0")
m("Core — Saetze",
  lambda w: f'=SUMIFS(tblLog[Sets],tblLog[Category],"Core",tblLog[Week],{w})',
  '=SUMIFS(tblLog[Sets],tblLog[Category],"Core")', "0.0")
m("Ø TopRPE",
  lambda w: f'=IFERROR(AVERAGEIFS(tblLog[TopRPE],tblLog[Week],{w}),0)',
  '=IFERROR(AVERAGE(tblLog[TopRPE]),0)', "0.0")
m("Ø Pain",
  lambda w: f'=IFERROR(AVERAGEIFS(tblLog[Pain_0_10],tblLog[Week],{w}),0)',
  '=IFERROR(AVERAGE(tblLog[Pain_0_10]),0)', "0.0")
m("Pain-Flags ≥4",
  lambda w: f'=COUNTIFS(tblLog[Pain_0_10],">=4",tblLog[Week],{w})',
  '=COUNTIF(tblLog[Pain_0_10],">=4")', "0")
m("Geloggte Zeilen",
  lambda w: f'=COUNTIFS(tblLog[Week],{w})',
  '=COUNTA(tblLog[Date])', "0")

rr = SUM_HR + 1
for label, wtmpl, gtmpl, numfmt in metrics:
    put(ws_sum, rr, 1, label, bold=True, align=LEFT)
    for wi, w in enumerate([1,2,3,4]):
        f_ = wtmpl(w) if callable(wtmpl) else wtmpl
        put(ws_sum, rr, 2+wi, f_, align=CENTER, numfmt=numfmt)
    put(ws_sum, rr, 6, gtmpl, align=CENTER, numfmt=numfmt, bold=True)
    rr += 1
SUM_LAST = rr - 1

for col,w in {"A":40,"B":10,"C":10,"D":10,"E":10,"F":14}.items():
    ws_sum.column_dimensions[col].width = w
ws_sum.freeze_panes = "B4"

# ============================================================================
# 4) B6_Dashboard
# ============================================================================
dash = wb.create_sheet("B6_Dashboard")
title(dash, "A1", "B6 — Performance-Block Dashboard (Hand Balancing / Calisthenics)", "A1:F1")

# Block-Variable (eine Zelle, ganz oben) -> benannter Bereich BlockCode
put(dash, 3, 1, "Block =", bold=True, align=Alignment(horizontal="right", vertical="center"))
bc = put(dash, 3, 2, "B6", bold=True, align=CENTER, fillhex=C_BLOCK, size=12)
wb.defined_names.add(DefinedName("BlockCode", attr_text="B6_Dashboard!$B$3"))
put(dash, 3, 3, "Folgeblock: diese Zelle auf B7 setzen, Sheets klonen, Prefix anpassen.",
    align=LEFT); dash.merge_cells("C3:F3")

# Situationsbewertung
section(dash, "A5", "Situationsbewertung", "A5:F5")
situ = [
 "Fortgeschrittener, autodidaktischer Hand-Balancer (1,90 m, ~80,8 kg). Skills von Canes auf den Boden uebertragen.",
 "Erster Performance-Block nach 6-Monats-Diaet (90 -> ~81 kg). Maintenance, KEIN Defizit, begrenzte Recovery.",
 "OAHS: links stabil, rechts limitiert (Schulter-Mobility/Koordination). Unilaterale Arbeit leicht zugunsten rechts.",
 "90 Push-up: 1 saubere Rep meist gut, 2. Rep sehr schwer. Ziel = 3 saubere Reps am Stueck.",
 "Ruecken (a) LWS/sakral: alte Reizung nach RDL/Deadlift, am Ausheilen -> Hinge-Rebuild ueber Pain-Gates.",
 "Ruecken (b) BWS: Einklemm-Gefuehl mittig/oben, ohne Ausstrahlen -> Mobility-/Tissue-Baustein, monitoren, KEINE Diagnose.",
 "Prinzip: Skill-Qualitaet + Gelenk-/Rueckengesundheit schlagen Volumen und Fatigue-PRs.",
]
rr = 6
for s in situ:
    put(dash, rr, 1, "•", align=CENTER); dash.merge_cells(f"B{rr}:F{rr}")
    put(dash, rr, 2, s, align=LEFT)
    rr += 1
situ_end = rr - 1

# Zielhierarchie
rr += 1
section(dash, f"A{rr}", "Zielhierarchie (skill-first; keine Wettkaempfe -> Progress-Maximierung, Gesundheit als Constraint)", f"A{rr}:F{rr}")
goals = [
 "1) 90 Push-up -> 3 saubere Reps am Stueck (Hauptziel).",
 "2) OAHS Shapes + Flag-Pfad (Straddle -> Shapes; supported Pre-Flag -> Sideline-Hold -> reduzierter Support -> seltene freie Versuche).",
 "3) Modern-Handstand-Linie und weitere Shapes (Flieger/Figa offen, Fernziel, nicht erzwingen).",
 "4) Kraftunterbau (Planche/Push/Pull) als Mittel zum Zweck fuer 1-3.",
 "5) Hypertrophie nur soweit sie Skills oder Schultergesundheit stuetzt.",
]
rr += 1
for g in goals:
    dash.merge_cells(f"A{rr}:F{rr}")
    put(dash, rr, 1, g, align=LEFT)
    rr += 1

# Wochen-Themen / Volumen-Leitplanken
rr += 1
section(dash, f"A{rr}", "Wochen-Themen & Volumen-Leitplanken (effektive harte Saetze/Woche)", f"A{rr}:I{rr}")
rr += 1
vol_headers = ["Woche / Thema","Push*","Pull/Upper-Back","Beine","90-Touches","OAHS-Bloecke","Planche-Dosen"]
header_row(dash, rr, vol_headers)
vol_hr = rr
VOL = [
 ("W1 Re-Entry / Baseline","12-14","14-16","8-10","3 (1 schwer / 2 leicht)","3","2"),
 ("W2 Build","14-16","16-18","10-12","3","3","2-3"),
 ("W3 Peak","15-16","17-18","12-14","3 (1 schwer / 2 assist)","3","3"),
 ("W4 Consolidate / Test","9-12","10-12","6-8","2 (Quality-Test)","beste Holds","1"),
]
rr += 1
for row in VOL:
    for i,val in enumerate(row):
        put(dash, rr, 1+i, val, align=LEFT if i==0 else CENTER)
    rr += 1
put(dash, rr, 1, "*Push zaehlt 90, Planche, Dips, HSPU, Vertikalpress zusammen. Taegliche Skill-Mikrodosis (6x/Woche) zaehlt NICHT ins harte Volumen.",
    align=LEFT); dash.merge_cells(f"A{rr}:G{rr}"); dash.cell(rr,1).font=F(8, italic=True)
rr += 2

# Live-Metriken
section(dash, f"A{rr}", "Live-Metriken (Formeln aus Log / Weekly_Summary)", f"A{rr}:E{rr}")
rr += 1
live_hr = rr
header_row(dash, rr, ["Metrik","Wert","Target","Interpretation","Action"])
rr += 1
# (Label, Wert-Formel, Target-Text, Interpretation-Formel, Action-Text, numfmt)
B = lambda a: f"$B${a}"  # nicht genutzt
def live(label, valf, target, interpf, action, numfmt="0.0"):
    global rr
    put(dash, rr, 1, label, bold=True, align=LEFT)
    vcell = f"B{rr}"
    put(dash, rr, 2, valf, align=CENTER, numfmt=numfmt)
    put(dash, rr, 3, target, align=CENTER)
    put(dash, rr, 4, interpf.format(v=vcell), align=LEFT)
    put(dash, rr, 5, action, align=LEFT)
    rr += 1

live("Geloggte Zeilen gesamt", "=COUNTA(tblLog[Date])", "wachsend",
     '=IF({v}=0,"noch leer",IF({v}<10,"wenig Daten","laeuft"))',
     "Taeglich 1 Zeile/Uebung anhaengen.", "0")
live("Ø TopRPE (Block)", "=IFERROR(AVERAGE(tblLog[TopRPE]),0)", "7,0-8,5",
     '=IF({v}=0,"-",IF({v}>8.7,"zu hoch",IF({v}<6.5,"Luft nach oben","im Ziel")))',
     "Wenn dauerhaft >8,5: Akzessorik kuerzen, Skill frisch halten.")
live("Ø Pain (Block)", "=IFERROR(AVERAGE(tblLog[Pain_0_10]),0)", "<2,5",
     '=IF({v}<2.5,"OK",IF({v}<3.5,"beobachten","regredieren"))',
     "Pain >3 oder Morgen schlechter -> Stufe runter (Rules).")
live("Pain-Flags ≥4", "=COUNTIF(tblLog[Pain_0_10],\">=4\")", "0",
     '=IF({v}=0,"OK","STOP/abklaeren")',
     "≥4 / scharf / neurologisch -> Uebung stoppen, ggf. abklaeren.", "0")
live("Saubere 90-Reps gesamt", '=SUMIFS(tblLog[CleanReps],tblLog[ExerciseID],"90")', "Trend hoch",
     '=IF({v}=0,"-","Fortschritt loggen")',
     "Ziel: Rep 1 submaximal -> Weg zu 3 sauberen Reps.", "0")
live("OAHS Best-Hold links (s)", '=MAXIFS(tblLog[BestHold_s],tblLog[Category],"SkillOAHS",tblLog[Side],"L")', "Trend hoch",
     '=IF({v}=0,"-",IF({v}>=5,"stabil","aufbauen"))',
     "Nur saubere Holds zaehlen.", "0")
live("OAHS Best-Hold rechts (s)", '=MAXIFS(tblLog[BestHold_s],tblLog[Category],"SkillOAHS",tblLog[Side],"R")', "naeher an links",
     '=IF({v}=0,"-",IF({v}>=MAXIFS(tblLog[BestHold_s],tblLog[Category],"SkillOAHS",tblLog[Side],"L")*0.7,"Asymmetrie ok","rechts priorisieren"))',
     "Rechts mehr Reps/Holds, Overhead/BWS-Mobility.", "0")
live("Flag Best-Hold links (s)", '=MAXIFS(tblLog[BestHold_s],tblLog[Category],"SkillFlag",tblLog[Side],"L")', "Trend hoch",
     '=IF({v}=0,"-","Support reduzieren wenn sauber")',
     "Support nur reduzieren bei sauberem Re-Entry.", "0")
live("Flag Best-Hold rechts (s)", '=MAXIFS(tblLog[BestHold_s],tblLog[Category],"SkillFlag",tblLog[Side],"R")', "Trend hoch",
     '=IF({v}=0,"-","Support reduzieren wenn sauber")',
     "Kontrollierter Return Pflicht.", "0")

dash_live_last = rr - 1
for col,w in {"A":34,"B":12,"C":16,"D":26,"E":40,"F":10,"G":14,"H":14,"I":14}.items():
    dash.column_dimensions[col].width = w

# CF auf Live Pain/RPE Wertspalte (B) – Pain-Flags & Pain Zeilen
dash.conditional_formatting.add(f"B{live_hr+3}", CellIsRule(operator="greaterThanOrEqual", formula=["2.5"], fill=fill(C_ORANGE)))
dash.conditional_formatting.add(f"B{live_hr+4}", CellIsRule(operator="greaterThanOrEqual", formula=["1"], fill=fill(C_RED)))

# ============================================================================
# 5) B6_Schedule
# ============================================================================
sched = wb.create_sheet("B6_Schedule")
title(sched, "A1", "B6 — Schedule (4 Wochen x 5 Tage + taegliche Skill-Mikrodosis)", "A1:J1")
sched["A2"] = ("Jede Einheit startet mit Skill-Mikrodosis (8-15 min): Handgelenke -> HS-Line/Balance -> OAHS-Line/Weight-Shift. "
               "Keine Max-Versuche an Nicht-Skill-Tagen. Skill nie ans Ende ermuedeter Sessions.")
sched["A2"].font = F(9, italic=True); sched.merge_cells("A2:J2"); sched["A2"].alignment = WRAP

S_HR = 3
s_headers = ["Woche","Tag","Day Type","Main Focus","Skill-Touchpoint","Strength","Back/Hinge-Stress","Placement","Progression-Target","Notes"]
header_row(sched, S_HR, s_headers)

# Tages-Grundgeruest (gilt fuer alle Wochen, Progression-Target variiert)
DAY_BASE = {
 "T1": ("90 Heavy + OAHS Line + Vertikalkraft",
        "OAHS Line/Weight-Shift + 2-3 saubere Holds/Seite",
        "90 Cluster (Singles->Double) + Pseudo-Planche/Planche-Hold + Weighted Pull-up + Scap/AR",
        "niedrig (kein Hinge)", "nach Ruhe-/Low-Tag (hoechster CNS-Tag)"),
 "T2": ("OAHS Shapes + Pull-Volumen",
        "OAHS Shape-Density (Straddle->Half/Diamond), nur saubere Holds",
        "Chest-supported/Seal Row + Pull-up/Pulldown + Rear-Delt/Face Pull + Hollow/Compression",
        "niedrig (kein LWS-Reiz)", "Pull-Schwerpunkt, Skill frisch"),
 "T3": ("Beine (wirbelsaeulen-freundlich) + Core + BWS-Mobility",
        "nur Line, keine Max-Versuche",
        "BSS/FFE Split Squat + Leg Press/Hack/Belt + Ham Curl + Calves + Pallof/Dead Bug/Side Plank",
        "Gate-gesteuert: Hinge mit Dowel -> entscheidet den Tag; KEIN schwerer Hinge",
        "BWS-Mobility-Block fix (Extension/Rotation, ribs down, Atmung)"),
 "T4": ("OAHS Flag/Transitions + 90 Light/Assisted",
        "Laengeres Warm-up + Side-Body-Aktivierung",
        "Supported Pre-Flag-Lean + Sideline/Mini-Flag + 90 leicht/Negativ+Pause + Vertikalpress + OA/Chest Row",
        "niedrig (ribs down, kein LWS-Arch)", "Support nur reduzieren wenn Re-Entry sauber"),
 "T5": ("Upper Hypertrophy (getrimmt) + 3. 90-Mikrodosis",
        "90 Technik/Negativ-Mikro (low volume, sauber)",
        "Push-Pump (Ring-Dips/Maschine/Deficit) + Pull-Pump (Row+Pulldown/Chin) + Schultern + Arme opt.",
        "niedrig", "BWS-Mobility + Lat/Pec/Forearm + Atmung-Downshift"),
}
WEEK_THEME = {1:"Re-Entry/Baseline",2:"Build",3:"Peak",4:"Consolidate/Test"}
WEEK_PROG = {
 1: {"T1":"Baseline: 90 4 Saetze x1, RPE<=8; Holds Qualitaet finden",
     "T2":"Shape-Baseline, saubere Straddle-Holds zaehlen",
     "T3":"Hinge-Gate testen (Level), Lasten konservativ",
     "T4":"Support hoch, Flag-Geometrie lernen",
     "T5":"Pump leicht, Technik 90 Negativ"},
 2: {"T1":"90 5 Saetze, Cluster 1->Double anpeilen, RPE 8-8,5",
     "T2":"Shape-Density +1 Block, Übergang Half-Straddle",
     "T3":"Hinge ggf. Level +1, Beinlast +",
     "T4":"Support leicht reduzieren wenn sauber",
     "T5":"Pump +1-2 Saetze, 90 Negativ +Pause"},
 3: {"T1":"90 Peak: saubere Double in 1-2 Saetzen, RPE<=8,5 (kein Failure)",
     "T2":"beste Shapes konsolidieren, Diamond/Legs-together-Prep",
     "T3":"Beinlast Peak, Hinge halten (nicht steigern wenn Pain)",
     "T4":"Mini-Flag-Hold laenger, seltene reduzierte Supports",
     "T5":"Pump Peak getrimmt, 90 Technik scharf"},
 4: {"T1":"Quality-Test: saubere Reps/Holds beweisen, Volumen runter",
     "T2":"beste Holds testen, kein Grind",
     "T3":"Deload Beine, Mobility-Fokus",
     "T4":"beste Flag-Holds dokumentieren",
     "T5":"Mini-Dosis, Regeneration"},
}
rr = S_HR + 1
for w in [1,2,3,4]:
    for dt in ["T1","T2","T3","T4","T5"]:
        focus, skill, strength, back, place = DAY_BASE[dt]
        put(sched, rr, 1, f"W{w} {WEEK_THEME[w]}", align=CENTER)
        put(sched, rr, 2, dt, align=CENTER, bold=True)
        put(sched, rr, 3, focus.split(" + ")[0] if False else dt, align=CENTER)
        put(sched, rr, 4, focus, align=LEFT)
        put(sched, rr, 5, skill, align=LEFT)
        put(sched, rr, 6, strength, align=LEFT)
        put(sched, rr, 7, back, align=LEFT)
        put(sched, rr, 8, place, align=LEFT)
        put(sched, rr, 9, WEEK_PROG[w][dt], align=LEFT)
        put(sched, rr, 10, "Skill-Mikrodosis zuerst; bei Qualitaetsverlust 90 -> erst Push-Akzessorik kuerzen", align=LEFT)
        if w % 2 == 0:
            for c in range(1,11):
                if sched.cell(rr,c).fill.fgColor.rgb in (None,"00000000"):
                    sched.cell(rr,c).fill = fill(C_BAND)
        rr += 1
for col,w in {"A":18,"B":7,"C":10,"D":30,"E":34,"F":40,"G":34,"H":26,"I":34,"J":34}.items():
    sched.column_dimensions[col].width = w
sched.freeze_panes = "A4"

# ============================================================================
# 6) B6_Plan_Detail
# ============================================================================
plan = wb.create_sheet("B6_Plan_Detail")
title(plan, "A1", "B6 — Plan Detail (Uebungsbloecke je Tag T1-T5)", "A1:L1")
plan["A2"] = "W1-W4 = Saetze x Reps/Holds je Woche. RPE/Quality und Gates beachten. Skill immer zuerst, nie ermuedet."
plan["A2"].font = F(9, italic=True); plan.merge_cells("A2:L2"); plan["A2"].alignment = WRAP

P_HR = 3
p_headers = ["Day Type","Block","Uebung/Drill","W1","W2","W3","W4","RPE/Quality","Rest","Progression/Gate","Back-Rule","Notes"]
header_row(plan, P_HR, p_headers)

# (DayType, Block, Uebung, W1, W2, W3, W4, RPE/Quality, Rest, Progression/Gate, Back-Rule, Notes)
PLAN = [
 # Taeglich
 ("ALLE","Warm-up","Wrist Prep + Handgelenks-Kraeftigung","1 Block","1","1","1","Qual 4-5","-","fix vor jeder Einheit","neutral","wegen hoher Handstand-Last"),
 ("ALLE","Skill-Mikrodosis","HS Line/Balance","8-15 min","kurz","kurz","kurz","Qual 4-5","-","Frequenz > Intensitaet","ribs down","keine Max-Versuche an Nicht-Skill-Tagen"),
 ("ALLE","Skill-Mikrodosis","OAHS Line / Weight-Shift","kurz","kurz","kurz","kurz","Qual 4-5","-","taeglich, nur sauber","neutral","rechts leicht bevorzugen"),
 # T1
 ("T1","Skill","OAHS Line + 2-3 saubere Holds/Seite","3/Seite","3","3","2 (Test)","Qual 4-5","60-90s","nur saubere Holds zaehlen","neutral","rechts ggf. +1 Hold"),
 ("T1","T1 Haupt","90 Heavy — Cluster Singles->Double","4x1","5x1","2x2 / Rest 1","Quality-Test","RPE 8-8,5","2-3 min","immer 1 Rep VOR Formverlust stoppen; nie Failure default","kein Arch, ribs down","Weg zu 3: Kraftdach + Cluster + Hebel"),
 ("T1","Planche","Pseudo-Planche Push-up ODER Adv-Tuck/Straddle Hold","3x6-8 / 3x8-12s","3","3","1","RPE 7-8","2 min","Straight-Arm-Fokus; Lean/Hold steigern","neutral","Schluessel-Assistance fuer 90"),
 ("T1","Vertikal-Pull","Weighted Pull-up","4x4","4-5x3-5","5x3","3x3","RPE 7-8","2-3 min","Last +, Reps sauber","neutral","Zug-Kraftbasis"),
 ("T1","Schulter-Hygiene","Scap/Serratus + External Rotation","2x10-15","2","2","1","RPE 6","45-60s","Qualitaet, kein Ego","neutral","Schultergesundheit"),
 # T2
 ("T2","Skill","OAHS Shape-Density (Straddle->Half/Diamond)","3 Bloecke","3","3","beste Holds","Qual 4-5","60-90s","nur saubere Holds; Progression der Form","neutral","rechts priorisieren"),
 ("T2","Haupt-Zug","Chest-supported / Seal Row","3x10-12","4x10-12","4x8-10","2x10","RPE 7-8","90s","Last/Reps +","brustgestuetzt, kein LWS-Reiz","oberer Ruecken"),
 ("T2","Vertikal-Zug","Pull-up / Lat Pulldown Volumen","3x8-10","4x8-10","4x6-8","2x8","RPE 7-8","90s","Volumen steuern","neutral","Lat-Volumen"),
 ("T2","Schulter","Rear-Delt Fly + Face Pull","2-3x12-15","3","3","2","RPE 6-7","45-60s","Qualitaet","neutral","hintere Schulter/Rotatoren"),
 ("T2","Core","Hollow / Compression","3x20-30s","3","3","2","Qual 4","45s","Spannung halten","KEIN LWS-Reiz","Compression fuer Shapes"),
 # T3
 ("T3","Status/Gate","Hinge-Pattern mit Dowel (Status-Check)","2-3 Reps","2-3","2-3","2-3","Pain-Gate","-","Gate entscheidet Tag (Rules)","Pain >3 -> regredieren","LWS + BWS Status"),
 ("T3","Bein uni","Bulgarian / FFE Split Squat","3x8-10/Seite","3x8-10","4x8","2x10","RPE 7-8","90s","Last +, Tiefe sauber","wirbelsaeulen-neutral","Quad/Glute"),
 ("T3","Bein bi","Leg Press / Hack / Belt Squat","3x10-12","3x10-12","4x10","2x12","RPE 7-8","2 min","Last +","kein axialer Grind","Quad/Glute"),
 ("T3","Post. Kette","Hamstring Curl + (gated) leichte post. Kette","3x10-12","3","3","2","RPE 7","75s","Hip-Thrust nur schmerzfrei; KEIN schwerer Hinge","Gate","Hamstrings/Glute"),
 ("T3","Bein","Calves","3x12-15","3","3","2","RPE 7","45s","Last/Reps +","neutral","Waden"),
 ("T3","BWS-Mobility","BWS Extension/Rotation, ribs down, Atmung","1 Block","1","1","1","Qual 4-5","-","fix; Protokoll in Rules","Einklemm-Thema monitoren","kein Ausstrahlen = ok"),
 ("T3","Core","Pallof + Dead Bug + Side Plank","2-3 je","3","3","2","Qual 4","45-60s","Anti-Ext/Anti-Rot/Anti-Lat","Anti-Flexion","LWS-schonend"),
 # T4
 ("T4","Warm-up","Laenger + Side-Body-Aktivierung","1 Block","1","1","1","Qual 4-5","-","vor Flag noetig","neutral","Verletzungsschutz"),
 ("T4","Flag","Supported Pre-Flag-Lean","4 Entries/Seite","4","4-5","beste","Qual 4-5","90s","Support nur reduzieren wenn Re-Entry sauber","kontrolliert","Flag-Pfad"),
 ("T4","Flag","Sideline / Mini-Flag-Hold + kontrollierter Return","3/Seite","3-4","4","beste","Qual 4-5","90s","Hold-Zeit/Support progressiv","kontrollierter Return","seltene freie Versuche spaeter"),
 ("T4","90 Light","90 leicht/assistiert ODER Negativ+Pause","3-4x2-3","4x2-3","4x2","2x2","RPE 6-7","2 min","sauberes Volumen, KEIN Failure","ribs down","2. Wochendosis 90"),
 ("T4","Vertikal-Press","Wall-HSPU / Pike / DB","3x5-8","3x6-8","4x6","2x6","RPE 7-8","90s","ribs down, kein LWS-Arch","kein Arch","Vertikalkraft"),
 ("T4","Zug-Balance","One-Arm / Chest-supported Row","3x10/Seite","3","3","2","RPE 7","75s","rechts ggf. +","gestuetzt","Zug-Balance"),
 # T5
 ("T5","90 Mikro","90 Technik/Negativ-Mikro","3x2","3x2-3","3x2","2x2","RPE 6-7 / Qual 4-5","2 min","3. Wochendosis, low volume, sauber","neutral","Technik scharf halten"),
 ("T5","Push-Pump","Ring-Dips / Maschine / Deficit Push-up","3x10-12","3-4x10-12","4x10","2x12","RPE 7-8","75s","getrimmt; bei Skill-Abfall zuerst kuerzen","neutral","Push-Hypertrophie"),
 ("T5","Pull-Pump","Row + Pulldown/Chin","3x10-12","4x10-12","4x10","2x12","RPE 7-8","75s","komplettiert Wochen-Zug","neutral","Pull-Hypertrophie"),
 ("T5","Schulter","Lateral Raise + Rear-Delt","2-3x12-15","3","3","2","RPE 6-7","45s","getrimmt","neutral","Schultern"),
 ("T5","Arme opt.","Biceps / Triceps","2x10-15","2-3","2","1","RPE 7","45s","optional, nur wenn Recovery ok","neutral","Arme"),
 ("T5","Regen","BWS-Mobility + Lat/Pec/Forearm + Atmung-Downshift","1 Block","1","1","1","Qual 4-5","-","Downshift Parasympathikus","BWS-Pflege","Recovery"),
]
rr = P_HR + 1
cur_day = None
for row in PLAN:
    for i,val in enumerate(row):
        align = CENTER if i in (0,3,4,5,6,7,8) else LEFT
        put(plan, rr, 1+i, val, align=align)
    # Farbband je Day Type
    dtc = row[0]
    band = {"ALLE":"E2EFDA","T1":"FCE4D6","T2":"DDEBF7","T3":"FFF2CC","T4":"E2DFF5","T5":"F2F2F2"}.get(dtc,"FFFFFF")
    plan.cell(rr,1).fill = fill(band)
    rr += 1
for col,w in {"A":8,"B":14,"C":40,"D":16,"E":12,"F":14,"G":14,"H":14,"I":9,"J":40,"K":22,"L":30}.items():
    plan.column_dimensions[col].width = w
plan.freeze_panes = "C4"

# ============================================================================
# 7) B6_Nutrition_Targets
# ============================================================================
nut = wb.create_sheet("B6_Nutrition_Targets")
title(nut, "A1", "B6 — Nutrition Targets (Maintenance, KEIN Defizit)", "A1:G1")

section(nut, "A3", "Eingaben / Konstanten", "A3:C3")
put(nut, 4, 1, "Aktuelles Gewicht (kg)", bold=True); put(nut, 4, 2, 80.8, align=CENTER, numfmt="0.0")
wb.defined_names.add(DefinedName("BW", attr_text="B6_Nutrition_Targets!$B$4"))
put(nut, 5, 1, "Zielkorridor (kg)", bold=True); put(nut, 5, 2, "80,5 - 81,0", align=CENTER)
put(nut, 6, 1, "Protein g/kg", bold=True); put(nut, 6, 2, 2.2, align=CENTER, numfmt="0.0")
put(nut, 7, 1, "Fett g/kg (Minimum)", bold=True); put(nut, 7, 2, 0.9, align=CENTER, numfmt="0.0")
put(nut, 8, 1, "Maintenance Basis (kcal)", bold=True); put(nut, 8, 2, 2700, align=CENTER, numfmt="0")
put(nut, 9, 1, "Protein-Ziel (g) [Formel]", bold=True); put(nut, 9, 2, "=ROUND(BW*B6,0)", align=CENTER, numfmt="0")
put(nut, 10, 1, "Fett-Minimum (g) [Formel]", bold=True); put(nut, 10, 2, "=ROUND(BW*B7,0)", align=CENTER, numfmt="0")
put(nut, 4, 3, "CF: ausserhalb 80,5-81,0 wird rot markiert.", align=LEFT); nut.merge_cells("C4:G4")
put(nut, 6, 3, "Protein 2,1-2,3 g/kg (~170-185 g). Fett >=0,8 g/kg. Rest = Carbs um Training.", align=LEFT); nut.merge_cells("C6:G6")

# Day-Type Tabelle mit Formeln
section(nut, "A12", "Day-Types: kcal & Makros (Carbs per Formel = Rest)", "A12:G12")
N_HR = 13
n_headers = ["Day Type","kcal-Delta","kcal-Ziel","Protein (g)","Fett (g)","Carbs (g)","Hinweis"]
header_row(nut, N_HR, n_headers)
# (DayType, delta, hinweis)
NUT_DAYS = [
 ("T1 (CNS/heavy)", 250, "Mini-Refeed-Option: mehr Carbs vor schwerem Skill-Tag"),
 ("T2", 150, "Pull-Volumen, moderate Carbs"),
 ("T3 (Beine)", 200, "Beinlast, Carbs hoch"),
 ("T4 (Flag)", 150, "Skill-lastig, Carbs moderat-hoch"),
 ("T5 (Pump)", 100, "Hypertrophie getrimmt"),
 ("Rest", -150, "Ruhetag, Carbs runter, Protein halten"),
]
rr = N_HR + 1
for dt, delta, hint in NUT_DAYS:
    put(nut, rr, 1, dt, bold=True, align=LEFT)
    put(nut, rr, 2, delta, align=CENTER, numfmt="+0;-0")
    put(nut, rr, 3, f"=$B$8+B{rr}", align=CENTER, numfmt="0")   # kcal-Ziel
    put(nut, rr, 4, "=$B$9", align=CENTER, numfmt="0")           # Protein g
    put(nut, rr, 5, "=$B$10", align=CENTER, numfmt="0")          # Fett g
    put(nut, rr, 6, f"=ROUND((C{rr}-D{rr}*4-E{rr}*9)/4,0)", align=CENTER, numfmt="0")  # Carbs Rest
    put(nut, rr, 7, hint, align=LEFT)
    rr += 1

section(nut, f"A{rr+1}", "Anpass-Trigger", f"A{rr+1}:G{rr+1}")
rr += 2
TRIG = [
 "BW < 80,2 kg UND Performance runter  ->  +100 bis +150 kcal (zuerst Carbs).",
 "BW > 82 kg schnell UND Taille schlechter  ->  -100 bis -150 kcal.",
 "Vor schweren Skill-Tagen (T1/T4): optionaler Mini-Refeed (Carbs +).",
 "KEIN Defizit in diesem Block — Ziel ist Performance auf Maintenance.",
]
for t in TRIG:
    nut.merge_cells(f"A{rr}:G{rr}")
    put(nut, rr, 1, t, align=LEFT)
    rr += 1

for col,w in {"A":24,"B":12,"C":12,"D":12,"E":10,"F":10,"G":46}.items():
    nut.column_dimensions[col].width = w
# CF Gewicht ausserhalb Korridor
nut.conditional_formatting.add("B4", CellIsRule(operator="notBetween", formula=["80.5","81.0"],
        fill=fill(C_RED), font=F(10,bold=True,color=C_REDT)))

# ============================================================================
# 8) B6_Rules
# ============================================================================
rules = wb.create_sheet("B6_Rules")
title(rules, "A1", "B6 — Rules (Autoregulation, Gates, Hinge-Rebuild, BWS-Protokoll, Referenzen)", "A1:E1")

rr = 3
section(rules, f"A{rr}", "Vier-Stufen-Gates (Basis: Pain 0-10, Quality 0-5, RPE)", f"A{rr}:E{rr}")
rr += 1
header_row(rules, rr, ["Bereich","GREEN","YELLOW","ORANGE","RED"])
rr += 1
GATES = [
 ("General",
  "Pain 0-2, Quality 4-5, RPE planmaessig -> Last/Volumen leicht hoch",
  "Pain 3 ODER Quality 3 -> Last halten, Technik priorisieren",
  "Pain 3 + Morgen schlechter -> 1 Stufe regredieren, Volumen kuerzen",
  "Pain >=4 / scharf / neurologisch -> Uebung STOP, ggf. abklaeren"),
 ("90 Push-up",
  "Rep 1 sauber, kein Schmerz -> Cluster/Last steigern",
  "2. Rep Formverlust -> bei Singles bleiben, Hebel/Kraftdach",
  "Schulter sperrt/Beschwerden -> auf Negativ+Pause zurueck",
  "Schmerz Schulter/Ellbogen scharf -> stoppen"),
 ("OAHS / Flag",
  "saubere Holds, Return kontrolliert -> Support reduzieren",
  "wackelig/rechts sperrt -> Support halten, mehr Reps rechts",
  "Re-Entry unsauber -> Support erhoehen, Geometrie ueben",
  "Schulter scharf / Kontrollverlust -> stoppen"),
 ("Back / Hinge",
  "Pain 0-1, Pattern sauber -> Hinge-Level +1 (Rules-Leiter)",
  "Pain 2 -> Level halten, Volumen niedrig",
  "Pain 3 ODER Morgen schlechter -> Level -1, kein Hinge-Load",
  "Pain >=4 / Ausstrahlen / Taubheit -> STOP, abklaeren"),
]
for g in GATES:
    put(rules, rr, 1, g[0], bold=True, align=LEFT)
    put(rules, rr, 2, g[1], align=LEFT, fillhex=C_GREEN)
    put(rules, rr, 3, g[2], align=LEFT, fillhex=C_YELLOW)
    put(rules, rr, 4, g[3], align=LEFT, fillhex=C_ORANGE)
    put(rules, rr, 5, g[4], align=LEFT, fillhex=C_RED)
    rr += 1

rr += 1
section(rules, f"A{rr}", "Skill-Progressionsregeln", f"A{rr}:E{rr}")
rr += 1
SKILLR = [
 "Hoch nur bei GREEN ueber min. 2 Sessions: Last/Hold/Support eine Stufe.",
 "Runter sofort bei ORANGE/RED oder wenn naechster Morgen schlechter.",
 "90: Weg zu 3 Reps = Kraftdach anheben (Rep 1 submax) + Cluster-Strength-Endurance + sauberer Hebel. Nie Failure default.",
 "OAHS Flag-Pfad: supported Pre-Flag -> Sideline-Hold -> reduzierter Support -> seltene freie Versuche. Support nur senken wenn Re-Entry sauber.",
 "Schwaechere (rechte) Seite leicht bevorzugen: mehr Reps/Holds, Overhead-/BWS-Mobility.",
 "Skill immer zuerst und frisch; bei Qualitaetsverlust ZUERST Push-Akzessorik kuerzen, nicht den Skill.",
]
for s in SKILLR:
    rules.merge_cells(f"A{rr}:E{rr}")
    put(rules, rr, 1, "• " + s, align=LEFT)
    rr += 1

rr += 1
section(rules, f"A{rr}", "Hinge-Rebuild-Leiter (Level 0-3) — KEIN schwerer / Barbell-Hinge in diesem Block", f"A{rr}:E{rr}")
rr += 1
header_row(rules, rr, ["Level","Allowed when","Exercises","Prescription","Move up when"])
rr += 1
HINGE = [
 ("0","Pain bei Dowel-Hinge 0-1, ADL schmerzfrei","Dowel Hip-Hinge, Hip-Airplane leicht, Back-Ext Iso (kurz)","2-3x5-8 Pattern, Iso 2-3x10-20s, taeglich moeglich","2 Sessions GREEN, Morgen ok"),
 ("1","Level 0 GREEN stabil","Hip-Thrust leicht, 45° Back-Ext (BW), Bird-Dog","2-3x8-12, RPE<=6, schmerzfrei","2 Sessions GREEN, kein Naechst-Morgen-Reiz"),
 ("2","Level 1 GREEN","Hip-Thrust moderat, KB Deadlift LEICHT (Pattern), Ham Curl betont","3x8-10, RPE 6-7","2 Sessions GREEN"),
 ("3","Level 2 GREEN (Ende Block / B7)","Trap-Bar / RDL LEICHT-moderat (erst naechster Block)","2-3x6-8, RPE<=7 — NICHT in B6","Naechster Block, schmerzfrei"),
]
for h in HINGE:
    put(rules, rr, 1, h[0], bold=True, align=CENTER)
    for i in range(1,5):
        put(rules, rr, 1+i, h[i], align=LEFT)
    rr += 1
rules.merge_cells(f"A{rr}:E{rr}")
put(rules, rr, 1, "WICHTIG: In B6 maximal Level 2. Kein schwerer Barbell-Hinge, keine Max-Grinds als Default. Posteriore Kette ueber Ham-Curl, Back-Ext-Iso, optional Hip-Thrust (nur schmerzfrei).",
    align=LEFT); rules.cell(rr,1).font=F(9,bold=True,color=C_REDT)
rr += 2

section(rules, f"A{rr}", "BWS-Mobility-Protokoll (Einklemm-Thema, KEINE Diagnose — nur Training/Mobility)", f"A{rr}:E{rr}")
rr += 1
BWS = [
 "Extension: Foam-Roller BWS-Extension 2x8-10; Quadruped T-Spine Extension.",
 "Rotation: Open-Book / Quadruped Rotation 2x6-8/Seite, langsam.",
 "Rib-Position: 'ribs down', Bauch-Atmung im 90/90, 3-4 Atemzuege; kein Flaring unter Overhead-Last.",
 "Atmung/Downshift: 5 min nasale Ausatem-betonte Atmung am Sessionende (Parasympathikus).",
 "Platzierung: voll an T3 & T5, Kurzversion im taeglichen Warm-up. Monitoren via Pain-Spalte im Log.",
]
for b in BWS:
    rules.merge_cells(f"A{rr}:E{rr}")
    put(rules, rr, 1, "• " + b, align=LEFT)
    rr += 1

rr += 1
section(rules, f"A{rr}", "Referenzen (Quelle | URL | wofuer)", f"A{rr}:E{rr}")
rr += 1
header_row(rules, rr, ["Quelle","URL","Wofuer","",""])
rr += 1
REFS = [
 ("Schoenfeld u.a. (PMC8884877)","https://pmc.ncbi.nlm.nih.gov/articles/PMC8884877/","Hypertrophie-Volumen ~12-20 Saetze/Muskel (individuell)"),
 ("Load/Reps Kraft vs. Hypertrophie (PMC7927075)","https://pmc.ncbi.nlm.nih.gov/articles/PMC7927075/","Last-/Rep-Bereiche Kraft vs. Hypertrophie"),
 ("Handstand Factory","https://handstandfactory.com/one-arm-shapes/","OAHS One-Arm-Shapes (Prereqs, fortgeschritten)"),
 ("Berg Movement","https://www.bergmovement.com/calisthenics-blog/one-arm-handstand-drills-and-progressions-beginner","OAHS Drills/Progressionen"),
 ("Berg Movement","https://www.bergmovement.com/calisthenics-blog/90-degree-push-up-tutorial","90 Push-up Tutorial/Definition"),
 ("YouTube","https://www.youtube.com/watch?v=hLKoKIAp6Eg","90 Progressionen (Negativ/Band/Momentum)"),
 ("The Barbell Physio","https://thebarbellphysio.com/returning-to-deadlifts-after-back-pain/","Return-to-Deadlift nach Rueckenschmerz (Hinge zuerst)"),
]
for src,url,wof in REFS:
    put(rules, rr, 1, src, align=LEFT)
    c = put(rules, rr, 2, url, align=LEFT); c.font = F(9, color="0563C1")
    put(rules, rr, 3, wof, align=LEFT); rules.merge_cells(f"C{rr}:E{rr}")
    rr += 1

for col,w in {"A":26,"B":40,"C":34,"D":30,"E":30}.items():
    rules.column_dimensions[col].width = w

# ============================================================================
# 9) Log_Guide
# ============================================================================
guide = wb.create_sheet("Log_Guide")
title(guide, "A1", "Log_Guide — Bedien- & Voice-Anleitung", "A1:D1")

rr = 3
section(guide, f"A{rr}", "Schema B6_Training_Log (genaue Reihenfolge)", f"A{rr}:D{rr}")
rr += 1
header_row(guide, rr, ["Spalte","Typ / Validierung","Bedeutung",""])
guide.merge_cells(f"C{rr}:D{rr}")
rr += 1
SCHEMA = [
 ("Date","echtes Datum","Trainingstag ('heute' = aktuelles Datum)"),
 ("DayType","Dropdown T1-T5, Rest","Session-Typ"),
 ("Week","1-4","Mesozyklus-Woche"),
 ("ExerciseID","Dropdown aus Library","Uebung/Skill"),
 ("Category","FORMEL (XLOOKUP)","Auto-Kategorie — NICHT manuell"),
 ("Side","Dropdown L,R,NA","Seite (OAHS/unilateral), sonst NA"),
 ("Sets","Zahl","Anzahl Saetze/Bloecke"),
 ("Reps_or_Scheme","Text","z.B. 8 oder 5/3/1 oder Holds-Anzahl"),
 ("Load_kg","Zahl","Zusatzlast (BW-Uebungen = 0)"),
 ("TopRPE","0-10 (0,5er)","hoechste RPE des Eintrags"),
 ("Pain_0_10","0-10","Schmerz waehrend/dazu"),
 ("Quality_0_5","0-5","Skill-Qualitaet (bei Skill)"),
 ("BestHold_s","Zahl","bester Hold in Sekunden (bei Holds)"),
 ("CleanReps","Zahl","saubere Reps (v.a. 90/Skill)"),
 ("Notes","Text","freie Voice-Notiz"),
 ("Key","FORMEL =DayType&\"_W\"&Week","fuer Rollups — NICHT manuell"),
]
for sp,ty,be in SCHEMA:
    put(guide, rr, 1, sp, bold=True, align=LEFT)
    put(guide, rr, 2, ty, align=LEFT)
    put(guide, rr, 3, be, align=LEFT); guide.merge_cells(f"C{rr}:D{rr}")
    rr += 1

rr += 1
section(guide, f"A{rr}", "Regeln fuer das Logging-Modell (Klartext)", f"A{rr}:D{rr}")
rr += 1
GRULES = [
 "Append-only: neue Eintraege IMMER unten an B6_Training_Log anhaengen, nichts ueberschreiben.",
 "Alias -> ExerciseID ueber B6_Exercise_Library mappen. Bei nicht eindeutigem Begriff EINE kurze Rueckfrage, sonst nicht nachfragen.",
 "Nicht zutreffende Felder = NA oder leer (z.B. BestHold_s bei einer Kraftuebung).",
 "Category und Key werden per Formel gefuellt (Tabellen-Spaltenformel zieht automatisch) — NICHT manuell eintragen.",
 "Datum als echtes Datum; 'heute' = aktuelles Datum.",
 "Eine Zeile pro Uebung pro Session. OAHS/Flag mit beiden Seiten -> zwei Zeilen (Side=L und Side=R).",
]
for g in GRULES:
    guide.merge_cells(f"A{rr}:D{rr}")
    put(guide, rr, 1, "• " + g, align=LEFT)
    rr += 1

rr += 1
section(guide, f"A{rr}", "Beispiel-Diktate -> exakte Zeile(n)", f"A{rr}:D{rr}")
rr += 1
header_row(guide, rr, ["#","Diktat","Resultierende Zeile(n)",""])
guide.merge_cells(f"C{rr}:D{rr}")
rr += 1
EX = [
 ("1","Heute T1, Woche 2. 90 Grad, 5 Saetze, je 1 saubere Wiederholung, RPE 8,5, Ruecken 1, Handgelenk gut.",
  "Date=heute, DayType=T1, Week=2, ExerciseID=90, Side=NA, Sets=5, Reps_or_Scheme=1, Load_kg=0, TopRPE=8.5, Pain_0_10=1, CleanReps=5, Notes=Handgelenk gut  (Category/Key auto)"),
 ("2","OAHS Holds, links 3 Holds bester 7 Sekunden Qualitaet 4, rechts 2 Holds bester 4 Sekunden Qualitaet 3.",
  "ZWEI Zeilen, ExerciseID=OAHS_line: (Side=L, Sets=3, BestHold_s=7, Quality_0_5=4)  und  (Side=R, Sets=2, BestHold_s=4, Quality_0_5=3)"),
 ("3","Weighted Pull-up 4 Saetze 4 Wiederholungen plus 20 Kilo RPE 7.",
  "ExerciseID=WPU, Category(auto)=Pull, Side=NA, Sets=4, Reps_or_Scheme=4, Load_kg=20, TopRPE=7"),
 ("4","Supported Flag links 4 Entries Qualitaet 4 Ruecken 2.",
  "ExerciseID=OAHS_flag_supported, Side=L, Sets=4, Quality_0_5=4, Pain_0_10=2"),
 ("5","Pseudo Planche Push-up 3 mal 8 RPE 8.",
  "ExerciseID=pseudo_planche_pushup, Category(auto)=Planche, Sets=3, Reps_or_Scheme=8, TopRPE=8"),
 ("6","T3 Woche 1, Hinge mit Dowel 3 Reps Ruecken 1, dann Beinpresse 3 mal 12 mit 120 Kilo RPE 7.",
  "ZWEI Zeilen: (ExerciseID=hinge_pattern, Sets=3, Reps_or_Scheme=3, Pain_0_10=1)  und  (ExerciseID=leg_press, Sets=3, Reps_or_Scheme=12, Load_kg=120, TopRPE=7)"),
]
for n,dik,res in EX:
    put(guide, rr, 1, n, bold=True, align=CENTER)
    put(guide, rr, 2, dik, align=LEFT)
    put(guide, rr, 3, res, align=LEFT); guide.merge_cells(f"C{rr}:D{rr}")
    rr += 1

rr += 1
section(guide, f"A{rr}", "Kurze Bedienanleitung", f"A{rr}:D{rr}")
rr += 1
HOWTO = [
 "1) Taeglich: pro Uebung EINE Zeile unten in B6_Training_Log anhaengen (Dropdowns nutzen). Skill zuerst loggen.",
 "2) Voice: diktiere DayType, Woche, Uebung (Alias reicht), Saetze/Reps/Last, RPE, Pain, ggf. Quality/Hold/Seite. Das Modell mappt Alias->ExerciseID und haengt an.",
 "3) Category & Key fuellen sich automatisch (Tabellenformel). Niemals ueberschreiben.",
 "4) B6_Weekly_Summary und B6_Dashboard aktualisieren sich automatisch (Pain/RPE-Ampeln, Volumen, Best-Holds).",
 "5) PLAN anpassen: in B6_Plan_Detail / B6_Schedule (Saetze/Reps/Progression). TRACKING anpassen: nur in B6_Training_Log.",
 "6) Gates & Hinge-Leiter in B6_Rules befolgen. Ernaehrung in B6_Nutrition_Targets (Maintenance).",
 "7) Folgeblock B7: B6_-Sheets klonen, BlockCode (B6_Dashboard!B3) auf B7 setzen, Prefix anpassen.",
]
for h in HOWTO:
    guide.merge_cells(f"A{rr}:D{rr}")
    put(guide, rr, 1, h, align=LEFT)
    rr += 1

for col,w in {"A":8,"B":52,"C":40,"D":40}.items():
    guide.column_dimensions[col].width = w
guide.freeze_panes = "A2"

# Reihenfolge der Sheets: Dashboard zuerst
order = ["B6_Dashboard","B6_Schedule","B6_Plan_Detail","B6_Exercise_Library",
         "B6_Training_Log","B6_Nutrition_Targets","B6_Rules","B6_Weekly_Summary","Log_Guide"]
wb._sheets.sort(key=lambda s: order.index(s.title))
wb.active = wb.sheetnames.index("B6_Dashboard")

# Recalc beim Oeffnen erzwingen
wb.calculation.fullCalcOnLoad = True

# Sicherheits-Pass: alle Zellen auf Arial setzen (Groesse/Bold beibehalten)
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and c.font is not None and c.font.name != FONT:
                c.font = Font(name=FONT, size=c.font.size, bold=c.font.bold,
                              italic=c.font.italic, color=c.font.color)

wb.save("/home/user/My-Test-repository/Training_2.xlsx")
print("OK: Training_2.xlsx geschrieben.")
print("Sheets:", wb.sheetnames)
