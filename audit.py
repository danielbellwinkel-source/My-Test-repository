# -*- coding: utf-8 -*-
"""Statische Pruefung aller Formeln in Training_2.xlsx (ohne LibreOffice)."""
import openpyxl, re, sys

wb = openpyxl.load_workbook("Training_2.xlsx")
LOG_COLS = ["Date","DayType","Week","ExerciseID","Category","Side","Sets","Reps_or_Scheme",
            "Load_kg","TopRPE","Pain_0_10","Quality_0_5","BestHold_s","CleanReps","Notes","Key"]
LIB_COLS = ["ExerciseID","Aliases","DisplayName","Category","MuscleTag","TracksLoad","TracksHold","TracksSide"]
TABLES = {"tblLog": LOG_COLS, "tblLib": LIB_COLS}

problems = []
formula_count = 0
all_formulas = []

# 1) Strukturierte Referenzen pruefen: jedes tblX[col] muss existieren
ref_re = re.compile(r"(tbl\w+)\[(@?\[?#?[^\]\[]+?)\]")
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and v.startswith("="):
                formula_count += 1
                all_formulas.append((ws.title, c.coordinate, v))
                for m in ref_re.finditer(v):
                    tbl, col = m.group(1), m.group(2)
                    col = col.lstrip("@")
                    if col.startswith("#") or col.startswith("["):
                        continue  # special items / nested
                    if tbl not in TABLES:
                        problems.append(f"{ws.title}!{c.coordinate}: unbekannte Tabelle {tbl}")
                    elif col not in TABLES[tbl]:
                        problems.append(f"{ws.title}!{c.coordinate}: Spalte '{col}' nicht in {tbl}: {v}")

# 2) calculatedColumnFormula in tblLog
log = wb["B6_Training_Log"]
tdef = log.tables["tblLog"]
calc_cols = {tc.name: (tc.calculatedColumnFormula.attr_text if tc.calculatedColumnFormula else None)
             for tc in tdef.tableColumns}
print("calculatedColumnFormula Category:", calc_cols.get("Category"))
print("calculatedColumnFormula Key     :", calc_cols.get("Key"))
for name in ("Category","Key"):
    if not calc_cols.get(name):
        problems.append(f"tblLog: calculatedColumnFormula fuer {name} fehlt")

# 3) Defined names
print("\nDefined names:")
for n, dn in wb.defined_names.items():
    print(f"  {n} -> {dn.value}")

# 4) Datenvalidierungen
print("\nDataValidations B6_Training_Log:")
for dv in log.data_validations.dataValidation:
    print(f"  type={dv.type} formula1={dv.formula1} sqref={dv.sqref}")

# 5) Demo-Werte unabhaengig nachrechnen
lib = wb["B6_Exercise_Library"]
lib_map = {}
for r in range(4, lib.max_row+1):
    eid = lib.cell(r,1).value
    cat = lib.cell(r,4).value
    if eid: lib_map[eid] = cat
print("\nLibrary IDs:", len(lib_map))
# 90 -> Skill90?
print("XLOOKUP('90') -> Category:", lib_map.get("90"), "(erwartet Skill90)")
if lib_map.get("90") != "Skill90":
    problems.append("Demo Category-Lookup fuer '90' != Skill90")
# Demo Key
dt = log.cell(2,2).value; wk = log.cell(2,3).value
print(f"Key demo = {dt}_W{wk} (erwartet T1_W1)")
# Weekly push W1 = sum sets where cat in {Skill90,Planche,Push} & week=1
# Demo: 90 -> Skill90, week1, sets5 -> 5
push_w1 = 0
for r in range(2, log.max_row+1):
    eid = log.cell(r,4).value
    if eid is None: continue
    cat = lib_map.get(eid)
    wk2 = log.cell(r,3).value
    sets = log.cell(r,7).value or 0
    if cat in ("Skill90","Planche","Push") and wk2==1:
        push_w1 += sets
print("Weekly Push W1 (erwartet 5):", push_w1)

# 6) Alle Library-Kategorien gueltig?
ALLOWED = {"Push","Pull","Legs","Skill90","SkillOAHS","SkillFlag","Planche","Core","Mobility","Wrist"}
badcat = {eid:cat for eid,cat in lib_map.items() if cat not in ALLOWED}
if badcat:
    problems.append(f"Ungueltige Kategorien: {badcat}")
print("Alle Kategorien gueltig:", not badcat)

# 7) §5-Abdeckung: Schluesseluebungen vorhanden?
needed = ["90","OAHS_line","OAHS_shapes","OAHS_flag_supported","flag_sideline","pseudo_planche_pushup",
          "planche_lean","planche_hold","WPU","pullup","lat_pulldown","chest_row","oa_row","rear_delt_fly",
          "face_pull","scap_serratus","ext_rotation","wall_hspu","pike_press","ring_dips","deficit_pushup",
          "lateral_raise","bss","leg_press","ham_curl","hip_thrust","back_ext_iso","calves","hinge_pattern",
          "pallof","dead_bug","side_plank","hollow","bws_mobility","wrist_prep","chin_up"]
missing = [n for n in needed if n not in lib_map]
print("Fehlende Schluessel-IDs:", missing or "keine")
if missing: problems.append(f"Fehlende IDs: {missing}")

print(f"\nFormeln gesamt: {formula_count}")
print("="*50)
if problems:
    print(f"PROBLEME ({len(problems)}):")
    for p in problems: print("  -", p)
    sys.exit(1)
else:
    print("=> STATISCHE PRUEFUNG OK, keine Probleme.")
